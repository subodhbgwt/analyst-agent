"""Builds the one-slide PPTX from company.json, analysis.json and thesis.json."""
import json
import sys
import tempfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DARK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x16, 0x57, 0x88)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)


def _text(slide, x, y, w, h, text, size=10, bold=False, color=DARK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Arial"
    r.font.color.rgb = color
    return box


def _bullets(slide, x, y, w, h, items, size=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = f"• {item}"
        r.font.size = Pt(size)
        r.font.name = "Arial"
        r.font.color.rgb = DARK
        p.space_after = Pt(4)


def _chart(years, revenue, ebitda, currency, path):
    fig, ax1 = plt.subplots(figsize=(4.4, 2.0), dpi=200)
    scale = 1e9 if max(revenue) > 5e9 else 1e6
    unit = "bn" if scale == 1e9 else "mm"
    ax1.bar(years, [r / scale for r in revenue], color="#165788", label="Revenue")
    ax1.set_ylabel(f"Revenue ({currency}{unit})", fontsize=8)
    ax2 = ax1.twinx()
    m = [e / r * 100 if (e and r) else None for e, r in zip(ebitda, revenue)]
    ax2.plot(years, m, color="#C55A11", marker="o", label="EBITDA margin")
    ax2.set_ylabel("EBITDA margin (%)", fontsize=8)
    for ax in (ax1, ax2):
        ax.tick_params(labelsize=8)
        ax.spines[["top"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def render(company: dict, analysis: dict, thesis: dict, out: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    name = company["name"]
    cur = company.get("currency") or ""
    _text(slide, 0.5, 0.25, 12.3, 0.5, f"{name}: company profile", 22, True, ACCENT)

    # Left column: overview + thesis
    _text(slide, 0.5, 1.0, 5.9, 0.3, "Business overview", 12, True)
    _bullets(slide, 0.5, 1.35, 5.9, 2.2, thesis["overview"])
    _text(slide, 0.5, 3.6, 5.9, 0.3, "Why interesting", 12, True)
    _bullets(slide, 0.5, 3.95, 5.9, 1.6, thesis["angle"])
    _text(slide, 0.5, 5.6, 5.9, 0.3, "Risks / open questions", 12, True)
    _bullets(slide, 0.5, 5.95, 5.9, 1.2, thesis["risks"])

    # Right column: chart + financials + valuation
    years = analysis["years"]
    with tempfile.TemporaryDirectory() as td:
        png = Path(td) / "chart.png"
        _chart(years, analysis["revenue"], analysis["ebitda"], cur, png)
        slide.shapes.add_picture(str(png), Inches(6.7), Inches(1.0), width=Inches(6.1))

    rows = len(years) + 1
    tbl = slide.shapes.add_table(5, rows, Inches(6.7), Inches(4.0),
                                 Inches(6.1), Inches(1.6)).table
    heads = [f"{cur}mm"] + years
    lines = [
        ("Revenue", [f"{v/1e6:,.0f}" if v else "–" for v in analysis["revenue"]]),
        ("growth %", [f"{v:+.1%}" if v is not None else "–" for v in analysis["revenue_growth"]]),
        ("EBITDA", [f"{v/1e6:,.0f}" if v else "–" for v in analysis["ebitda"]]),
        ("margin %", [f"{v:.1%}" if v is not None else "–" for v in analysis["ebitda_margin"]]),
    ]
    for c, head in enumerate(heads):
        cell = tbl.cell(0, c)
        cell.text = head
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
        cell.text_frame.paragraphs[0].runs[0].font.bold = True
    for r, (label, vals) in enumerate(lines, start=1):
        tbl.cell(r, 0).text = label
        for c, v in enumerate(vals, start=1):
            tbl.cell(r, c).text = v
        for c in range(rows):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    d = analysis.get("dcf")
    val_lines = []
    if company.get("market_cap"):
        val_lines.append(f"Market cap {cur} {company['market_cap']/1e9:,.1f}bn | "
                         f"EV {cur} {company['enterprise_value']/1e9:,.1f}bn")
    if analysis.get("ev_ebitda"):
        pe = f" | P/E {analysis['pe']:.1f}x" if analysis.get("pe") else ""
        val_lines.append(f"EV/EBITDA {analysis['ev_ebitda']:.1f}x{pe}")
    if d:
        val_lines.append(
            f"Quick DCF ({d['wacc']:.1%} WACC, {d['tgr']:.1%} TGR, growth fading "
            f"{d['start_growth']:.1%}→{d['tgr']:.1%}): implied EV {cur} "
            f"{d['implied_ev']/1e9:,.1f}bn ({d['upside']:+.0%} vs market)")
    _text(slide, 6.7, 5.85, 6.1, 0.3, "Valuation snapshot", 12, True)
    _bullets(slide, 6.7, 6.2, 6.1, 1.1, val_lines)

    _text(slide, 0.5, 7.15, 12.3, 0.25,
          "Source: Yahoo Finance market data. Summary written by an AI agent, check it before use. "
          "DCF is a screening heuristic, not investment advice.", 7, False, GRAY)

    prs.save(out)


if __name__ == "__main__":
    workdir = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    company = json.loads((workdir / "company.json").read_text(encoding="utf-8"))
    analysis = json.loads((workdir / "analysis.json").read_text(encoding="utf-8"))
    thesis = json.loads((workdir / "thesis.json").read_text(encoding="utf-8"))
    out = workdir / f"{company['ticker'].replace('.', '_')}_profile.pptx"
    render(company, analysis, thesis, out)
    print(f"written to {out}")
